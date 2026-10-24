"""Compare two ROC-AUC reproduction runs (Run 2 minus Run 1).

Place this script alongside ROC_AUC_run1.xlsx and ROC_AUC_run2.xlsx,
then run: python compare_reproduction_runs.py
"""
from pathlib import Path
import re
import pandas as pd

BASE = Path(__file__).resolve().parent
RUN_1 = BASE / 'ROC_AUC_run1.xlsx'
RUN_2 = BASE / 'ROC_AUC_run2.xlsx'
OUTPUT = BASE / 'ROC_AUC_run2_minus_run1_computed.xlsx'
OUTPUT_CSV = BASE / 'ROC_AUC_run2_minus_run1_computed.csv'
ALGORITHMS = ['ABOD', 'CBLOF', 'HBOS', 'IForest', 'KNN', 'LODA', 'LOF', 'OCSVM', 'FB', 'COPOD']


def read_run(path):
    df = pd.read_excel(path, dtype=str)
    df.columns = df.columns.str.strip()
    if 'Data' not in df.columns or any(c not in df.columns for c in ALGORITHMS):
        raise ValueError(f'{path.name}: expected Data and all 10 algorithm columns')
    df = df[['Data'] + ALGORITHMS].copy()
    df['Data'] = df['Data'].str.strip().str.lower()
    if df['Data'].duplicated().any():
        raise ValueError(f'{path.name}: duplicate dataset labels')
    for col in ALGORITHMS:
        # Handles screenshot-style cells such as "0.7688 (9)" or plain numeric scores.
        df[col] = pd.to_numeric(
            df[col].str.replace(r'\s*\(\d+\)\s*$', '', regex=True).str.replace(',', '.', regex=False),
            errors='raise'
        )
    return df.set_index('Data')


def main():
    first, second = read_run(RUN_1), read_run(RUN_2)
    if set(first.index) != set(second.index):
        raise ValueError('Dataset lists differ between the two runs')
    second = second.reindex(first.index)
    difference = (second - first).round(4)
    difference.index.name = 'Data'
    difference.to_csv(OUTPUT_CSV, float_format='%.4f')

    with pd.ExcelWriter(OUTPUT, engine='openpyxl') as writer:
        difference.to_excel(writer, sheet_name='Run 2 minus Run 1')
        ws = writer.sheets['Run 2 minus Run 1']
        ws.freeze_panes = 'B2'
        ws.column_dimensions['A'].width = 28
        for col in 'BCDEFGHIJK':
            ws.column_dimensions[col].width = 15
        from openpyxl.styles import Font, PatternFill
        for cell in ws[1]:
            cell.fill = PatternFill('solid', fgColor='1E3A5F')
            cell.font = Font(bold=True, color='FFFFFF')
        for row in ws.iter_rows(min_row=2, min_col=2):
            for cell in row:
                cell.number_format = '0.0000;-0.0000;0'
                if cell.value is not None and abs(cell.value) > 0:
                    cell.fill = PatternFill('solid', fgColor='FFF2CC')

    print('Differences = Run 2 - Run 1 (positive means Run 2 scored higher)')
    for col in ALGORITHMS:
        n = (difference.drop(index='avg', errors='ignore')[col].abs() > 0).sum()
        print(f'{col:8}: {n:2} / 30 datasets changed')
    print(f'Created {OUTPUT.name} and {OUTPUT_CSV.name}')


if __name__ == '__main__':
    main()
