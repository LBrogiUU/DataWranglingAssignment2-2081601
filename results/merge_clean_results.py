
import pandas as pd
from pathlib import Path
from openpyxl.styles import Font
from openpyxl.utils import get_column_letter


# Exact column order from the original COPOD paper
ALGORITHMS = [
    "ABOD", "CBLOF", "FB", "HBOS", "IForest",
    "KNN", "LODA", "LOF", "OCSVM", "COPOD"
]

COLUMNS_TO_REMOVE = [
    "# Samples", "# Dimensions", "Outlier Perc"
]


def ask_file(default):
    filename = input(f"CSV file [{default}]: ").strip() or default
    path = Path(filename)

    if not path.is_file():
        raise FileNotFoundError(f"File not found: {path}")

    return path


def load_and_clean(filename, source):
    df = pd.read_csv(filename, dtype=str)

    # Clean column names
    df.columns = df.columns.str.strip()

    # Remove unnecessary columns
    df = df.drop(columns=COLUMNS_TO_REMOVE, errors="ignore")

    # Rename COD to COPOD
    df = df.rename(columns={"COD": "COPOD"})

    # Check that all required columns exist
    required = ["Data"] + ALGORITHMS
    missing = [col for col in required if col not in df.columns]

    if missing:
        raise ValueError(f"{filename}: missing columns {missing}")

    # Select columns BY NAME in the paper's exact order
    # This ensures values remain under the correct algorithms
    df = df[required].copy()

    # Identify the dataset source
    df["Data"] = df["Data"].astype(str) + f" ({source})"

    # Convert algorithm scores to numeric values
    for algorithm in ALGORITHMS:
        df[algorithm] = pd.to_numeric(
            df[algorithm], errors="raise"
        )

    return df


def merge_results(mat_file, arff_file):
    mat = load_and_clean(mat_file, "mat")
    arff = load_and_clean(arff_file, "arff")

    # Combine MAT and ARFF datasets
    merged = pd.concat([mat, arff], ignore_index=True)

    # Calculate average scores across all 30 datasets
    avg = {"Data": "AVG"}

    for algorithm in ALGORITHMS:
        avg[algorithm] = merged[algorithm].mean()

    merged = pd.concat(
        [merged, pd.DataFrame([avg])],
        ignore_index=True
    )

    # Rank algorithms for each dataset
    # Higher score = better performance
    ranks = merged[ALGORITHMS].rank(
        axis=1,
        ascending=False,
        method="min"
    )

    # Format scores and append rankings
    for algorithm in ALGORITHMS:
        merged[algorithm] = [
            f"{score:.4f} ({int(rank)})"
            if pd.notna(score) and pd.notna(rank)
            else ""
            for score, rank in zip(
                merged[algorithm], ranks[algorithm]
            )
        ]

    # Ensure final column order matches the paper
    merged = merged[["Data"] + ALGORITHMS]

    return merged


def save_results(df, name):
    # Save CSV
    df.to_csv(f"{name}.csv", index=False)

    # Save Excel with first-place scores in bold
    with pd.ExcelWriter(
        f"{name}.xlsx", engine="openpyxl"
    ) as writer:

        df.to_excel(
            writer,
            index=False,
            sheet_name="Results"
        )

        ws = writer.sheets["Results"]

        # Bold first-place scores
        for row in ws.iter_rows(min_row=2):
            for cell in row:
                if (
                    isinstance(cell.value, str)
                    and cell.value.endswith("(1)")
                ):
                    cell.font = Font(bold=True)

        # Adjust column widths
        for column in ws.columns:
            max_length = max(
                len(str(cell.value))
                if cell.value is not None else 0
                for cell in column
            )

            ws.column_dimensions[
                get_column_letter(column[0].column)
            ].width = min(max_length + 3, 30)

    print(f"Saved: {name}.csv and {name}.xlsx")


def main():
    print("Enter the benchmark CSV paths.")
    print("Press Enter to use the default filename.\n")

    # Read MAT and ARFF results
    roc_mat = ask_file("roc_mat.csv")
    roc_arff = ask_file("roc_arff.csv")
    ap_mat = ask_file("ap_mat.csv")
    ap_arff = ask_file("ap_arff.csv")

    # Merge Average Precision results
    ap_merged = merge_results(ap_mat, ap_arff)
    save_results(ap_merged, "AP_merged")

    # Merge ROC-AUC results
    roc_merged = merge_results(roc_mat, roc_arff)
    save_results(roc_merged, "ROC_AUC_merged")

    print("\nDone! All results cleaned and merged successfully.")


if __name__ == "__main__":
    main()
